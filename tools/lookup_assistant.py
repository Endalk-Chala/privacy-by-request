#!/usr/bin/env python3
"""
lookup_assistant.py  --  session-time people-search assistant for the
"Privacy by Request" exposure study.

What it does (one participant per run, in the room, with the participant watching):
  1. asks for participant ID, name, city and state at the prompt (never on the
     command line, so nothing lands in shell history);
  2. opens each of the six people-search sites in a visible browser, one at a time;
  3. waits for you to press Enter once the results page has loaded (and after you
     have cleared any CAPTCHA by hand);
  4. captures the page text (and, if you enable it, a screenshot);
  5. proposes a 0/1/2 score for each of the seven fields from simple text rules;
  6. shows you the proposal and lets you accept or overwrite each score;
  7. appends 42 rows to a CSV in the exact column order of the Exposure_Grid sheet.

What it does NOT do: run without you, run on more than one person per launch,
store the name (only the search string you confirm, as the grid already requires),
or try to get past CAPTCHAs or paywalls.

Usage:
  python lookup_assistant.py                 # live run
  python lookup_assistant.py --dry-run       # fictional name, no browser, exercises the flow
  python lookup_assistant.py --manual        # type scores you noted in a session; no browser, no name
  python lookup_assistant.py --check-urls    # fictional name, opens each site so you can verify the URL patterns
  python lookup_assistant.py --screenshots   # also save PNGs to ./captures/<ID>/
  python lookup_assistant.py --append-xlsx exposure_study_data_collection.xlsx
                                             # append the rows straight into Exposure_Grid

Requires: pip install playwright openpyxl ; python -m playwright install chromium
"""

import argparse, csv, datetime as dt, os, re, sys, urllib.parse as up

SITES = ["Spokeo", "Whitepages", "TruePeopleSearch", "FastPeopleSearch", "BeenVerified", "Intelius"]
FIELDS = ["Current address", "Past addresses", "Phone", "Email", "Relatives", "Associates", "Court/public records"]
PAYWALLED = {"BeenVerified", "Intelius"}

SCORE_HELP = "0 = not shown, 1 = partial or ambiguous (namesake, truncated, teaser), 2 = clearly shown for this person"


def slug(s, sep="-"):
    return sep.join(re.sub(r"[^A-Za-z0-9 ]", "", s).split())


def build_url(site, first, last, city, state):
    """Best-effort search URLs. These change; check them before each wave of sessions."""
    f, l, c, st = first.strip(), last.strip(), city.strip(), state.strip().upper()
    if site == "TruePeopleSearch":
        return ("https://www.truepeoplesearch.com/results?" +
                up.urlencode({"name": f"{f} {l}", "citystatezip": f"{c}, {st}"}))
    if site == "FastPeopleSearch":
        return f"https://www.fastpeoplesearch.com/name/{slug(f)}-{slug(l)}_{slug(c)}-{st}".lower()
    if site == "Whitepages":
        return f"https://www.whitepages.com/name/{slug(f)}-{slug(l)}/{slug(c)}-{st}"
    if site == "Spokeo":
        return f"https://www.spokeo.com/{slug(f)}-{slug(l)}/{st}/{slug(c)}"
    if site == "BeenVerified":
        return ("https://www.beenverified.com/app/search/person?" +
                up.urlencode({"fname": f, "ln": l, "city": c, "state": st}))
    if site == "Intelius":
        return f"https://www.intelius.com/people-search/{slug(f)}-{slug(l)}/{slug(c)}-{st}"
    raise ValueError(site)


# ---- scoring rules -------------------------------------------------------------
ADDR = re.compile(r"\b\d{1,6}\s+[A-Za-z0-9.'\- ]{2,40}\b(St|Street|Ave|Avenue|Rd|Road|Dr|Drive|Ln|Lane|Blvd|Ct|Court|Way|Pl|Place|Cir|Trl|Pkwy|Hwy)\b", re.I)
PHONE = re.compile(r"\(?\b\d{3}\)?[-. ]\d{3}[-. ]\d{4}\b")
EMAIL = re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.]+\b")


def propose(site, text, last):
    """Return {field: (score, why)} from the captured page text."""
    t = text or ""
    low = t.lower()
    n_addr = len(set(m.group(0) for m in ADDR.finditer(t)))
    n_phone = len(set(PHONE.findall(t)))
    n_email = len(set(EMAIL.findall(t)))
    has_rel = "relative" in low or "family member" in low
    has_assoc = "associate" in low
    has_court = any(k in low for k in ["court record", "criminal record", "public record", "arrest", "bankrupt", "lien", "traffic"])
    namesake = low.count(last.lower()) > 3   # many hits -> likely several people with this surname

    out = {}
    if site in PAYWALLED:
        found = last.lower() in low and any(k in low for k in ["we found", "records found", "view report", "see results", "results for"])
        base = 1 if found else 0
        for fld in FIELDS:
            out[fld] = (base, "teaser only (paywalled): 1 = record exists, verify fields manually" if found else "no teaser match")
        return out

    def sc(cond, strong):
        if not cond:
            return 0
        return 1 if namesake else (2 if strong else 1)

    out["Current address"] = (sc(n_addr >= 1, n_addr >= 1), f"{n_addr} street address pattern(s) on page")
    out["Past addresses"] = (sc(n_addr >= 2, n_addr >= 3), f"{n_addr} distinct address pattern(s)")
    out["Phone"] = (sc(n_phone >= 1, n_phone >= 1), f"{n_phone} phone pattern(s)")
    out["Email"] = (sc(n_email >= 1, n_email >= 1), f"{n_email} email pattern(s)")
    out["Relatives"] = (sc(has_rel, has_rel), "relatives/family section present" if has_rel else "no relatives section")
    out["Associates"] = (sc(has_assoc, has_assoc), "associates section present" if has_assoc else "no associates section")
    out["Court/public records"] = (sc(has_court, False), "records keywords present" if has_court else "no records keywords")
    if namesake:
        for k in out:
            out[k] = (out[k][0], out[k][1] + "; many surname hits, likely namesakes, check which card is the participant")
    return out


# ---- interaction ---------------------------------------------------------------
def ask(prompt, default=None):
    s = input(f"{prompt}{' [' + str(default) + ']' if default is not None else ''}: ").strip()
    return s if s else default


def confirm_scores(site, proposal):
    """One line per site: Enter accepts all proposed scores; typing seven digits overrides them all;
    typing 'e' edits field by field."""
    print(f"\n--- {site} ---  ({SCORE_HELP})")
    for fld in FIELDS:
        p, why = proposal[fld]
        print(f"  {fld:22s} {p}   {why}")
    digits = "".join(str(proposal[f][0]) for f in FIELDS)
    final = {}
    while True:
        s = ask(f"  Enter = accept {digits} | seven digits (order: addr past phone email rel assoc court) | e = edit one by one", digits)
        if s == "e":
            for fld in FIELDS:
                p, why = proposal[fld]
                while True:
                    v = ask(f"    {fld:22s} proposed {p}  Enter to accept, or 0/1/2", str(p))
                    if v in ("0", "1", "2"):
                        final[fld] = int(v)
                        break
            break
        if len(s) == 7 and all(ch in "012" for ch in s):
            final = {fld: int(ch) for fld, ch in zip(FIELDS, s)}
            break
        print("    enter seven digits using only 0, 1 and 2, or 'e', or press Enter")
    note = ask("  note for this site (optional, no personal content)", "")
    return final, note


def run_live(pid, first, last, city, state, want_png, capdir):
    from playwright.sync_api import sync_playwright
    rows, today = [], dt.date.today().isoformat()
    search_string = "name + city"   # what the grid records; never the name itself
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        ctx = browser.new_context(viewport={"width": 1280, "height": 900})
        page = ctx.new_page()
        for site in SITES:
            url = build_url(site, first, last, city, state)
            print(f"\nOpening {site} (the search URL is not printed here, so the name stays out of the terminal)")
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=60000)
            except Exception:
                print("  could not load. Navigate manually in the browser window.")
            input("  When the results page is fully loaded (clear any CAPTCHA by hand), press Enter... ")
            try:
                text = page.inner_text("body")
            except Exception:
                text = ""
            if want_png:
                os.makedirs(capdir, exist_ok=True)
                png = os.path.join(capdir, f"{pid}_{site}_{today}.png")
                page.screenshot(path=png, full_page=True)
                print(f"  screenshot saved: {png}  (restricted tier: encrypt or delete per protocol)")
            final, note = confirm_scores(site, propose(site, text, last))
            for fld in FIELDS:
                rows.append([pid, site, fld, final[fld], search_string, today, note])
        browser.close()
    return rows


def check_urls():
    """Opens each site with a fictional name so you can confirm the URL pattern lands on a results page."""
    from playwright.sync_api import sync_playwright
    first, last, city, state = "Dawit", "Example", "Saint Paul", "MN"
    print("URL check with a fictional name (Dawit Example, Saint Paul, MN). Nothing is recorded.")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_context(viewport={"width": 1280, "height": 900}).new_page()
        for site in SITES:
            url = build_url(site, first, last, city, state)
            print(f"\n{site}: {url}")
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=60000)
            except Exception as e:
                print(f"  could not load ({e})")
            input("  Does this land on a search-results page (even an empty one)? Note it, then press Enter... ")
        browser.close()
    print("\nIf a site did not land on results, edit its line in build_url() or navigate by hand during sessions.")


def run_manual(pid):
    """Enter scores already noted during a session: one line of seven digits per site. No browser, no name."""
    print(f"Manual entry for {pid}. For each site type seven digits (order: addr past phone email rel assoc court), "
          f"or 'skip' if that site was not checked.  {SCORE_HELP}")
    rows, today = [], dt.date.today().isoformat()
    date = ask("Date of the session (YYYY-MM-DD)", today)
    for site in SITES:
        while True:
            s = ask(f"  {site:18s} seven digits or skip")
            if s == "skip":
                break
            if s and len(s) == 7 and all(ch in "012" for ch in s):
                note = ask("    note (optional, no personal content)", "")
                for fld, ch in zip(FIELDS, s):
                    rows.append([pid, site, fld, int(ch), "name + city", date, note])
                break
            print("    enter exactly seven digits using 0, 1 and 2, or 'skip'")
    return rows


def run_dry(pid):
    """Exercises the flow with a fictional person and canned page text; no network."""
    canned = {
        "Spokeo": "Results for Dawit Example in Saint Paul, MN. Age 41. 1234 Como Ave, Saint Paul MN. Relatives: ... Phone (651) 555-0147",
        "Whitepages": "Dawit Example, Saint Paul MN. Current address 1234 Como Ave. Past addresses: 55 Oak St, 980 Rice Street. (651) 555-0147. Relatives: Example, A. Associates",
        "TruePeopleSearch": "Dawit Example  Age 41  1234 Como Ave Saint Paul MN  Phone Numbers (651) 555-0147  Relatives  Associates  Email dawit@example.com",
        "FastPeopleSearch": "Dawit Example  1234 Como Ave  Saint Paul, MN  Past Addresses 55 Oak St  Relatives",
        "BeenVerified": "We found 1 records for Dawit Example in Saint Paul, MN. View report.",
        "Intelius": "No results for this search.",
    }
    rows, today = [], dt.date.today().isoformat()
    for site in SITES:
        final, note = confirm_scores(site, propose(site, canned[site], "Example"))
        for fld in FIELDS:
            rows.append([pid, site, fld, final[fld], "name + city", today, note or "DRY RUN"])
    return rows


HEADER = ["id", "site", "field", "score (0/1/2)", "search_string_used", "lookup_date", "notes (no content copied)"]


def write_csv(rows, path):
    """Append rows; if the file is locked (open in Excel), keep asking, and as a last resort
    save to a timestamped recovery file so a finished session is never lost."""
    while True:
        try:
            new = not os.path.exists(path)
            with open(path, "a", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                if new:
                    w.writerow(HEADER)
                w.writerows(rows)
            print(f"\n{len(rows)} rows appended to {path}")
            return
        except PermissionError:
            print(f"\n{path} is locked (probably open in Excel). Close it, then press Enter to retry, or type 's' to save to a recovery file.")
            if input("> ").strip().lower() == "s":
                alt = f"recovery_{rows[0][0]}_{dt.datetime.now():%Y%m%d_%H%M%S}.csv"
                with open(alt, "w", newline="", encoding="utf-8") as fh:
                    w = csv.writer(fh); w.writerow(HEADER); w.writerows(rows)
                print(f"saved to {alt}; merge it into {path} later")
                return


def append_xlsx(rows, path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Exposure_Grid"]
    # (workbook save is retried below if Excel has it open)
    # The template pre-lays out 42 rows per participant (id, site, field) with empty scores.
    # Fill those rows if they exist; otherwise append at the bottom.
    index = {}
    for r in range(2, ws.max_row + 1):
        key = (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)
        if key[0] not in (None, ""):
            index[key] = r
    filled, appended = 0, 0
    next_row = ws.max_row + 1
    for row in rows:
        key = tuple(row[:3])
        r = index.get(key)
        if r is None:
            r = next_row; next_row += 1; appended += 1
        else:
            filled += 1
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    while True:
        try:
            wb.save(path)
            break
        except PermissionError:
            print(f"{path} is locked (probably open in Excel). Close it, then press Enter to retry.")
            input("> ")
    print(f"Exposure_Grid in {path}: {filled} existing rows filled, {appended} rows appended")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dry-run", action="store_true", help="fictional person, no browser")
    ap.add_argument("--manual", action="store_true", help="type scores noted during a session; no browser, no name")
    ap.add_argument("--check-urls", action="store_true", help="open each site with a fictional name to verify URL patterns")
    ap.add_argument("--screenshots", action="store_true", help="save full-page PNGs to captures/<ID>/")
    ap.add_argument("--csv", default="exposure_grid_rows.csv", help="CSV to append rows to")
    ap.add_argument("--append-xlsx", metavar="WORKBOOK", help="also append rows into the Exposure_Grid sheet of this workbook")
    a = ap.parse_args()

    if a.check_urls:
        check_urls()
        return

    print("Lookup assistant. Run this only during a consented session, with the participant present.")
    pid = ask("Participant ID (e.g. P07)")
    if not pid:
        sys.exit("ID required")
    digits = re.sub(r"[^0-9]", "", pid.upper().replace("O", "0"))
    if digits:
        pid = f"P{int(digits):02d}"
    print(f"Recording as {pid}")
    if a.dry_run:
        rows = run_dry(pid)
    elif a.manual:
        rows = run_manual(pid)
        if not rows:
            sys.exit("nothing entered")
    else:
        first = ask("First name"); last = ask("Last name"); city = ask("City"); state = ask("State (2 letters)", "MN")
        if not all([first, last, city, state]):
            sys.exit("name, city and state are required")
        rows = run_live(pid, first, last, city, state, a.screenshots, os.path.join("captures", pid))
    write_csv(rows, a.csv)
    if a.append_xlsx:
        append_xlsx(rows, a.append_xlsx)
    print("Done. Review the scores against your session notes before analysis.")


if __name__ == "__main__":
    main()
